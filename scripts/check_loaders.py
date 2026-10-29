"""Smoke-test every supported ingestion format without touching the vector store.

Usage:
    python scripts/check_loaders.py
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.loaders import SUPPORTED_EXTENSIONS, load_file

results = []


def check(name, condition, detail=""):
    results.append((name, condition))
    print(f"{'PASS' if condition else 'FAIL'}  {name}" + (f"  {detail}" if not condition else ""))


def build_samples(root: Path):
    (root / "notes.txt").write_text("Plain text marker ALPHA\nsecond line", encoding="utf-8")
    (root / "guide.md").write_text("# Title\nMarkdown marker BRAVO", encoding="utf-8")
    (root / "page.html").write_text(
        "<html><head><title>Ops</title></head><body><h1>Runbook</h1>"
        "<p>HTML marker CHARLIE</p><script>ignore()</script></body></html>",
        encoding="utf-8",
    )
    (root / "config.json").write_text('{"key": "JSON marker DELTA"}', encoding="utf-8")
    (root / "table.csv").write_text("name,value\nalpha,CSV marker ECHO\n", encoding="utf-8")

    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("DOCX marker FOXTROT")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "k"
    table.cell(0, 1).text = "v"
    table.cell(1, 0).text = "cell"
    table.cell(1, 1).text = "DOCX table GOLF"
    doc.save(root / "report.docx")

    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Deck"
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "PPTX marker HOTEL"
    slide.notes_slide.notes_text_frame.text = "speaker note INDIA"
    prs.save(root / "deck.pptx")

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["item", "amount"])
    ws.append(["xlsx", "XLSX marker JULIET"])
    wb.create_sheet("Empty")
    wb.save(root / "sheet.xlsx")

    from reportlab.pdfgen import canvas

    pdf_path = root / "doc.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(100, 750, "PDF marker KILO")
    c.showPage()
    c.save()


def combined(docs):
    return "\n".join(d.page_content for d in docs)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        build_samples(root)

        files = {p.suffix for p in root.iterdir()}
        check("all sample types supported", files <= SUPPORTED_EXTENSIONS, str(files))

        expectations = {
            "notes.txt": "ALPHA",
            "guide.md": "BRAVO",
            "page.html": "CHARLIE",
            "config.json": "DELTA",
            "table.csv": "ECHO",
            "report.docx": "FOXTROT",
            "deck.pptx": "HOTEL",
            "sheet.xlsx": "JULIET",
            "doc.pdf": "KILO",
        }
        for name, marker in expectations.items():
            docs = load_file(root / name)
            ok = marker in combined(docs)
            check(f"{name} text extracted", ok)

        docx_docs = load_file(root / "report.docx")
        check("docx table extracted", "GOLF" in combined(docx_docs))

        pptx_docs = load_file(root / "deck.pptx")
        check("pptx notes extracted", "INDIA" in combined(pptx_docs))
        check(
            "pptx slide metadata",
            pptx_docs and pptx_docs[0].metadata.get("slide") == 1,
            str(pptx_docs[0].metadata if pptx_docs else None),
        )

        xlsx_docs = load_file(root / "sheet.xlsx")
        check(
            "xlsx sheet metadata",
            any(d.metadata.get("sheet") == "Budget" for d in xlsx_docs),
            str([d.metadata for d in xlsx_docs]),
        )

        html_docs = load_file(root / "page.html")
        check("html script stripped", "ignore()" not in combined(html_docs))
        check(
            "html title captured",
            html_docs and html_docs[0].metadata.get("title") == "Ops",
        )

        try:
            load_file(root / "notes.txt").pop()
            (root / "bad.bin").write_bytes(b"\x00\x01")
            load_file(root / "bad.bin")
            check("unsupported type rejected", False)
        except ValueError:
            check("unsupported type rejected", True)

    failed = [name for name, ok in results if not ok]
    print()
    if failed:
        print(f"RESULT: FAILED ({len(failed)}): {', '.join(failed)}")
        return 1
    print("RESULT: OK - all loaders passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
