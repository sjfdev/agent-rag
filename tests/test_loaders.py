"""Loaders: every supported format extracts its text and metadata."""

from pathlib import Path

import pytest

from src.loaders import SUPPORTED_EXTENSIONS, load_file


def _combined(docs):
    return "\n".join(doc.page_content for doc in docs)


@pytest.fixture()
def samples(tmp_path: Path):
    (tmp_path / "notes.txt").write_text("Plain ALPHA", encoding="utf-8")
    (tmp_path / "guide.md").write_text("# Title\nMarkdown BRAVO", encoding="utf-8")
    (tmp_path / "page.html").write_text(
        "<html><head><title>Ops</title></head><body>"
        "<p>HTML CHARLIE</p><script>ignore()</script></body></html>",
        encoding="utf-8",
    )
    (tmp_path / "config.json").write_text('{"k": "JSON DELTA"}', encoding="utf-8")
    (tmp_path / "table.csv").write_text("a,b\nx,CSV ECHO\n", encoding="utf-8")

    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("DOCX FOXTROT")
    table = doc.add_table(rows=2, cols=2)
    table.cell(1, 0).text = "cell"
    table.cell(1, 1).text = "DOCX table GOLF"
    doc.save(tmp_path / "report.docx")

    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Deck"
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "PPTX HOTEL"
    slide.notes_slide.notes_text_frame.text = "speaker INDIA"
    prs.save(tmp_path / "deck.pptx")

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["item", "amount"])
    ws.append(["xlsx", "XLSX JULIET"])
    wb.create_sheet("Empty")
    wb.save(tmp_path / "sheet.xlsx")

    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(tmp_path / "doc.pdf"))
    c.drawString(100, 750, "PDF KILO")
    c.showPage()
    c.save()
    return tmp_path


@pytest.mark.parametrize(
    "name, marker",
    [
        ("notes.txt", "ALPHA"),
        ("guide.md", "BRAVO"),
        ("page.html", "CHARLIE"),
        ("config.json", "DELTA"),
        ("table.csv", "ECHO"),
        ("report.docx", "FOXTROT"),
        ("deck.pptx", "HOTEL"),
        ("sheet.xlsx", "JULIET"),
        ("doc.pdf", "KILO"),
    ],
)
def test_each_format_extracts_text(samples, name, marker):
    docs = load_file(samples / name)
    assert marker in _combined(docs)


def test_supported_extensions_cover_samples(samples):
    suffixes = {path.suffix for path in samples.iterdir()}
    assert suffixes <= SUPPORTED_EXTENSIONS


def test_docx_tables_are_extracted(samples):
    assert "GOLF" in _combined(load_file(samples / "report.docx"))


def test_pptx_notes_and_slide_metadata(samples):
    docs = load_file(samples / "deck.pptx")
    assert "INDIA" in _combined(docs)
    assert docs[0].metadata.get("slide") == 1


def test_xlsx_sheet_metadata(samples):
    docs = load_file(samples / "sheet.xlsx")
    assert any(doc.metadata.get("sheet") == "Budget" for doc in docs)


def test_html_strips_scripts_and_keeps_title(samples):
    docs = load_file(samples / "page.html")
    assert "ignore()" not in _combined(docs)
    assert docs[0].metadata.get("title") == "Ops"


def test_unsupported_extension_rejected(samples):
    bad = samples / "bad.bin"
    bad.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError):
        load_file(bad)
